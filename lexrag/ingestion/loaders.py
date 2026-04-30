from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from lexrag.utils.logging import get_logger

logger = get_logger(__name__)


@dataclass
class LoadedDocument:
    doc_id: str
    file_path: str
    file_name: str
    file_type: str
    raw_text: str
    metadata: dict = field(default_factory=dict)


class DocumentLoader:
    """Load PDF, DOCX, and TXT files into raw text."""

    SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".text"}

    def load(self, file_path: str | Path) -> LoadedDocument:
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED:
            raise ValueError(f"Unsupported file type: {suffix}")

        doc_id = self._make_doc_id(path)
        logger.info(f"Loading {suffix} file: {path.name}")

        if suffix == ".pdf":
            text = self._load_pdf(path)
        elif suffix in {".docx", ".doc"}:
            text = self._load_docx(path)
        else:
            text = self._load_txt(path)

        text = self._normalize_text(text)
        return LoadedDocument(
            doc_id=doc_id,
            file_path=str(path),
            file_name=path.name,
            file_type=suffix.lstrip("."),
            raw_text=text,
            metadata={"source": path.name},
        )

    def _make_doc_id(self, path: Path) -> str:
        import hashlib
        return hashlib.md5(str(path).encode()).hexdigest()[:12]

    def _load_pdf(self, path: Path) -> str:
        try:
            import pdfplumber
            pages: list[str] = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages.append(text)
            return "\n\n".join(pages)
        except ImportError:
            pass

        # Fallback: pypdf
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages.append(text)
            return "\n\n".join(pages)
        except ImportError as e:
            raise RuntimeError("Install pdfplumber or pypdf to load PDFs") from e

    def _load_docx(self, path: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except ImportError as e:
            raise RuntimeError("Install python-docx to load DOCX files") from e

    def _load_txt(self, path: Path) -> str:
        for enc in ("utf-8", "latin-1", "cp1252"):
            try:
                return path.read_text(encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Cannot decode text file: {path}")

    def _normalize_text(self, text: str) -> str:
        # Collapse excessive whitespace while preserving paragraph breaks
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
